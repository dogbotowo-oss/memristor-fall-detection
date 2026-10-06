# -*- coding: utf-8 -*-
from pathlib import Path
import csv
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE = Path(r"F:\12.18-2\5.9new\crossbar+\CNN+SNN\v-8-final-new-0625")
REPORTS = BASE / "reports"
OUT_PATH = BASE / "v8-0625.docx"
IMAGE_PATH = BASE / "image" / "training_curve_v8final_new_0625.png"
V8_SUMMARY = Path(
    r"F:\12.18-2\5.9new\crossbar+\CNN+SNN\v8\reports"
    r"\snn_adl_gate_penalty_scan_event_norm\subject4_compare_p002_p006"
    r"\subject4_summary_p006.json"
)


def load_json(path):
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def pct(value):
    return f"{float(value) * 100:.2f}%"


def set_font(run, size=10.5, bold=False):
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)


def add_title(doc, title, subtitle):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    set_font(run, 18, True)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(subtitle)
    set_font(run, 12)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_font(run, 14 if level == 1 else 12, True)
    return paragraph


def add_para(doc, text=""):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    set_font(run, 9, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, header, True)
        shade(cell, "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    doc.add_paragraph("")
    return table


def video_ids(paths):
    ids = []
    for part in (paths or "").split(";"):
        part = part.strip()
        if part:
            ids.append(Path(part).stem)
    return ", ".join(ids)


def collect_text(doc):
    parts = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def main():
    subject4 = load_json(REPORTS / "v8final_new_0625_subject4_final_metrics.json")
    rec = load_json(REPORTS / "v8final_new_0625_recommended_threshold.json")
    dataset_compare = load_json(REPORTS / "dataset_compare_v8_vs_new0625.json")
    v8 = load_json(V8_SUMMARY)
    with (REPORTS / "v8final_new_0625_strict_validation_threshold_scan.csv").open(
        "r", encoding="utf-8-sig", newline=""
    ) as f:
        val_rows = list(csv.DictReader(f))

    best_val = max(
        val_rows,
        key=lambda row: (
            float(row["balanced_accuracy"]),
            float(row["adl_specificity"]),
            float(row["threshold"]),
        ),
    )

    if OUT_PATH.exists():
        OUT_PATH.unlink()

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    add_title(
        doc,
        "v8-0625 技术说明文档",
        "v-8-final-new-0625：基于 v8 final 的细粒度 ADL subtype 定向抑制实验",
    )
    add_para(doc, "生成日期：2026-06-26")
    add_para(doc, r"工程根目录：F:\12.18-2")
    add_para(doc, f"文档路径：{OUT_PATH}")
    add_para(
        doc,
        "说明：本文档根据已完成训练日志、严格验证结果、Subject4 最终测试 JSON/CSV "
        "和当前项目记录整理；本次不重新训练模型，且不设置打开密码。",
    )

    add_heading(doc, "1. 实验目的")
    add_para(
        doc,
        "本版本沿用 v8 final 的 CNN+SNN 主干、Crossbar / 忆阻器器件动态映射以及 "
        "Gaussian teacher-student noise，尝试路线 A：把统一 ADL penalty 改为更细粒度的 "
        "ADL subtype 定向抑制。",
    )
    add_bullets(
        doc,
        [
            "核心目标是降低坐下、弯腰、蹲跪等接近地面但未跌倒动作对 Fall 的误报。",
            "同时尽量保持 v8 final 在真实跌倒上的召回能力。",
            "所有阈值和策略选择仅允许在 Subject3 + external validation 上完成，Subject4 只作为独立最终测试。",
        ],
    )

    add_heading(doc, "2. 严格数据划分")
    add_table(
        doc,
        ["数据来源", "用途", "约束"],
        [
            ["GMDCSA Subject 1-2", "训练", "可用于模型参数学习"],
            ["GMDCSA Subject 3", "验证与阈值校准", "允许用于选择阈值和验证策略"],
            ["GMDCSA Subject 4", "独立最终测试", "禁止用于选模型、调阈值、调 sampler、调 penalty"],
            ["Zenodo / external data", "外部训练/验证补充", "必须按 video / person / subject 划分，禁止窗口随机混合"],
        ],
    )
    add_para(
        doc,
        f"阈值选择来源：{rec.get('source', 'Subject3 + Zenodo val')}。阈值选择策略："
        f"{rec.get('threshold_selection_policy', 'maximize balanced_accuracy; ties choose higher ADL specificity; remaining ties choose higher threshold')}。",
    )

    add_heading(doc, "3. 代码改动与新增机制")
    add_para(
        doc,
        r"主代码文件：F:\12.18-2\5.9new\crossbar+\CNN+SNN\v-8-final-new-0625\code\automatic_fall_detector.py",
    )
    add_bullets(
        doc,
        [
            "新增参数 --targeted-adl-subtype-weights，用于对不同 ADL subtype 赋予不同 penalty 权重。",
            "新增参数 --targeted-adl-penalty-hard-only，用于控制该 penalty 是否只作用于 hard negative ADL 窗口。",
            "新增参数 --targeted-adl-penalty-score-threshold，用于设置 hard negative score 生效阈值。",
            "新增 ADL subtype head，作为辅助学习分支，用于区分 sit_descent、bend_reach、squat_kneel 等 ADL 内部类型。",
            "保留 v8 final 的 CNN+SNN 主体、Crossbar/device dynamics、teacher-student Gaussian noise 和 SNN gate penalty 主线。",
        ],
    )
    add_table(
        doc,
        ["Subtype", "Penalty 权重", "设计意图"],
        [
            ["sit_descent", "0.04", "轻度抑制坐下/下降类 ADL 误报"],
            ["bend_reach", "0.05", "中等抑制弯腰/伸手接近地面动作"],
            ["squat_kneel", "0.07", "较强抑制蹲下/跪下类接近地面动作"],
        ],
    )
    add_para(
        doc,
        "生效条件：仅当 ADL 窗口的 hard_negative_score >= 0.30 时启用 targeted ADL penalty，"
        "避免普通 ADL 样本被过度压制。",
    )

    add_heading(doc, "4. 训练配置与数据规模")
    add_table(
        doc,
        ["项目", "配置/结果"],
        [
            [
                "初始化模型",
                r"F:\12.18-2\5.9new\crossbar+\CNN+SNN\v8\models\automatic_fall_event_detector_cnn_snn_v8_final_p006_penalty006.pth",
            ],
            ["Python 环境", r"F:\ANACONDA\envs\my_yizu_3.10\python.exe"],
            ["Crossbar 数据目录", r"F:\12.18-2\data"],
            ["external-labeled-stride", "16"],
            ["image-size", "64"],
            ["snn-adl-gate-penalty", "0.06"],
            ["hard-negative-normal-weight", "1.10"],
            ["adl-subtype-loss-weight", "0.10"],
            ["targeted-adl-penalty-score-threshold", "0.30"],
        ],
    )
    dataset_rows = []
    for item in dataset_compare:
        dataset_rows.append(
            [
                item["tag"],
                item["train_windows"],
                item["train_counts"].get("0"),
                item["train_counts"].get("3"),
                item["val_windows"],
                item["val_counts"].get("0"),
                item["val_counts"].get("3"),
                item["train_videos"],
                item["val_videos"],
            ]
        )
    add_table(
        doc,
        ["版本", "train windows", "train ADL", "train Fall", "val windows", "val ADL", "val Fall", "train videos", "val videos"],
        dataset_rows,
    )
    add_para(
        doc,
        "本次复现的窗口数为 train=2508、val=564，与当前 v8 路线一致；不要因 v14/v15 日志曾出现 "
        "3539/1076 就直接判断为数据错误。",
    )

    add_heading(doc, "5. 训练过程观察")
    add_bullets(
        doc,
        [
            "日志确认已加载 conductance data：HRS=9164, LRS=10001。",
            "日志确认已加载 device dynamics：EPSC、PPF、LTP、LTD。",
            "Epoch 03 达到较好的验证平衡点：val_acc=74.5%、val_bal=73.8%、val_normal=71.6%、val_fall=76.0%、val_subtype=41.1%。",
            "后续 epoch 出现 normal 与 fall 指标摆动，说明 subtype 信息被学习到，但没有稳定转化为更好的主分类边界。",
            "Epoch 08 early stopping，模型与训练曲线均已保存。",
        ],
    )
    if IMAGE_PATH.exists():
        add_para(doc, "训练曲线如下：")
        doc.add_picture(str(IMAGE_PATH), width=Inches(5.8))

    add_heading(doc, "6. 严格验证阈值扫描")
    add_para(
        doc,
        f"推荐阈值：{rec.get('threshold', 0.55)}。该阈值来自 "
        f"{rec.get('source', 'Subject3 + Zenodo val')}，不是 Subject4。",
    )
    add_table(
        doc,
        ["threshold", "Accuracy", "Balanced Acc", "ADL specificity", "Fall recall", "Precision", "TP/TN/FP/FN"],
        [
            [
                row["threshold"],
                pct(row["accuracy"]),
                pct(row["balanced_accuracy"]),
                pct(row["adl_specificity"]),
                pct(row["fall_recall"]),
                pct(row["precision"]),
                f"{row['tp']}/{row['tn']}/{row['fp']}/{row['fn']}",
            ]
            for row in val_rows
        ],
    )
    add_para(
        doc,
        f"最佳验证点：threshold={best_val['threshold']}，Accuracy={pct(best_val['accuracy'])}，"
        f"Balanced Accuracy={pct(best_val['balanced_accuracy'])}，ADL specificity={pct(best_val['adl_specificity'])}，"
        f"Fall recall={pct(best_val['fall_recall'])}，Precision={pct(best_val['precision'])}，"
        f"TP/TN/FP/FN={best_val['tp']}/{best_val['tn']}/{best_val['fp']}/{best_val['fn']}。",
    )

    add_heading(doc, "7. Subject4 独立最终测试")
    add_table(
        doc,
        ["指标", "v-8-final-new-0625"],
        [
            ["Threshold", subject4["threshold"]],
            ["Total videos", subject4["total"]],
            ["Accuracy", pct(subject4["accuracy"])],
            ["Balanced Accuracy", pct(subject4["balanced_accuracy"])],
            ["ADL specificity", pct(subject4["adl_specificity"])],
            ["Fall recall", pct(subject4["fall_recall"])],
            ["Precision", pct(subject4["precision"])],
            ["TP/TN/FP/FN", f"{subject4['tp']} / {subject4['tn']} / {subject4['fp']} / {subject4['fn']}"],
        ],
    )
    add_para(doc, "误报 ADL：" + video_ids(subject4.get("false_alarms", "")))
    add_para(doc, "漏检 Fall：" + video_ids(subject4.get("missed_falls", "")))

    add_heading(doc, "8. 与 v8 final p006 对比")
    add_table(
        doc,
        ["版本", "Threshold", "Accuracy", "Balanced Acc", "ADL specificity", "Fall recall", "Precision", "TP/TN/FP/FN"],
        [
            [
                "v8 final p006",
                v8["threshold"],
                pct(v8["accuracy"]),
                pct(v8["balanced_accuracy"]),
                pct(v8["adl_specificity"]),
                pct(v8["fall_recall"]),
                pct(v8["precision"]),
                f"{v8['tp']} / {v8['tn']} / {v8['fp']} / {v8['fn']}",
            ],
            [
                "v-8-final-new-0625",
                subject4["threshold"],
                pct(subject4["accuracy"]),
                pct(subject4["balanced_accuracy"]),
                pct(subject4["adl_specificity"]),
                pct(subject4["fall_recall"]),
                pct(subject4["precision"]),
                f"{subject4['tp']} / {subject4['tn']} / {subject4['fp']} / {subject4['fn']}",
            ],
        ],
    )
    add_para(
        doc,
        "对比结论：v-8-final-new-0625 没有超过 v8 final p006。它在 ADL specificity 上仍有一定控制能力，"
        "但 Fall recall 从 76.47% 降至 58.82%，Balanced Accuracy 从 78.24% 降至 66.91%。",
    )

    add_heading(doc, "9. 结论")
    add_bullets(
        doc,
        [
            "路线 A 的代码实现是成立的：按 ADL subtype 分权重、hard-only penalty、score threshold 均已跑通。",
            "当前参数组合使模型更保守，ADL 误报控制没有完全崩坏，但真实跌倒漏检明显增加。",
            "因此 v-8-final-new-0625 不能替代 v8 final p006 作为当前主线论文结果。",
            "本版本的价值主要是证明 subtype 信息可以接入训练流程，为后续 gate/post-score 校正或更轻量 subtype 抑制提供基础。",
        ],
    )

    add_heading(doc, "10. 后续建议")
    add_bullets(
        doc,
        [
            "保留当前 subtype head 和 targeted penalty 框架，但不要继续简单增大 penalty。",
            "优先尝试更轻量的 subtype 抑制，或只对最关键单一 subtype 做约束。",
            "更建议把 subtype 信息用于 SNN gate 或 post-score 校正，而不是直接压低主分类 Fall 概率。",
            "任何新阈值、sampler 权重、penalty 权重仍必须只在 Subject3 + external val 上选择，Subject4 继续作为一次性最终测试。",
        ],
    )

    add_heading(doc, "11. 文件记录")
    add_table(
        doc,
        ["类型", "路径"],
        [
            ["代码", str(BASE / "code" / "automatic_fall_detector.py")],
            ["模型", str(BASE / "models" / "automatic_fall_event_detector_cnn_snn_v8final_new_0625.pth")],
            ["训练日志", str(REPORTS / "train_v8final_new_0625.stdout.log")],
            ["训练曲线", str(IMAGE_PATH)],
            ["阈值扫描", str(REPORTS / "v8final_new_0625_strict_validation_threshold_scan.csv")],
            ["推荐阈值", str(REPORTS / "v8final_new_0625_recommended_threshold.json")],
            ["Subject4 结果", str(REPORTS / "v8final_new_0625_subject4_final_metrics.json")],
            ["Subject4 video metrics", str(REPORTS / "v8final_new_0625_subject4_final_video_metrics.csv")],
        ],
    )

    doc.save(str(OUT_PATH))

    checked = Document(str(OUT_PATH))
    text = collect_text(checked)
    bad_tokens = ["锛", "涓", "鐨", "绗", "杩", "浠", "鏄", "鍦", "濂", "℃", "????"]
    found_bad = {token: text.count(token) for token in bad_tokens if token in text}
    required = [
        "技术说明文档",
        "细粒度 ADL subtype 定向抑制实验",
        "严格数据划分",
        "对比结论",
        "不能替代 v8 final p006",
        "不设置打开密码",
    ]
    missing = [item for item in required if item not in text]
    if found_bad or missing:
        raise RuntimeError(f"文档内容检查失败 found_bad={found_bad}, missing={missing}")

    print("OK")
    print(str(OUT_PATH))
    print(OUT_PATH.stat().st_size)
    print(len(checked.paragraphs))
    print(text.splitlines()[0])
    print(text.splitlines()[1])


if __name__ == "__main__":
    main()

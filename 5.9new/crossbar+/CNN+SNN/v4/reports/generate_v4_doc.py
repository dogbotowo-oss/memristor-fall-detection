# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


V4_DIR = Path(r"F:\12.18-2\5.9new\crossbar+\CNN+SNN\v4")
OUT_PATH = V4_DIR / "v4-module.docx"


def set_font(run, font_name: str = "Microsoft YaHei", size: float | None = None) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size is not None:
        run.font.size = Pt(size)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_metric_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "指标"
    table.rows[0].cells[1].text = "结果"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value


def main() -> None:
    doc = Document()
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5 if style_name == "Normal" else 14)

    title = doc.add_heading("v4 模型说明文档：CNN + SNN 动态分支", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.add_run("工程路径：").bold = True
    p.add_run(str(V4_DIR))
    p = doc.add_paragraph()
    p.add_run("运行环境：").bold = True
    p.add_run(r"F:\ANACONDA\envs\my_yizu_3.10\python.exe")
    p = doc.add_paragraph()
    p.add_run("文档目的：").bold = True
    p.add_run("记录 v4 版本的 SNN 动态分支优化、训练/验证/测试流程、严格评估结果，以及当前问题和后续优化方向。")

    doc.add_heading("1. 工程背景", level=1)
    doc.add_paragraph(
        "本项目总体方向是神经网络 + Crossbar / 忆阻器器件表征，用视频跌倒检测作为任务场景。"
        "前期模型主要依赖 CNN 提取每帧空间特征，再通过 GRU 建模短时间视频窗口，最后进行 fall / ADL 判定。"
        "v4 的目标是在该框架上加入 LIF-inspired SNN 时间分支，让模型更关注跌倒瞬间的快速状态变化。"
    )
    add_bullets(
        doc,
        [
            "Crossbar 与器件噪声开关继续保留，用于模拟器件读出阶段的不确定性。",
            "hard negative ADL 逻辑继续保留，用于提高模型对疑似跌倒 ADL 的辨别能力。",
            "Subject4 严格只作为最终独立测试集，不参与阈值选择或模型调参。",
        ],
    )

    doc.add_heading("2. 数据划分与严格评估规则", level=1)
    add_bullets(
        doc,
        [
            "训练集：GMDCSA Subject 1-2 + Zenodo video split train。",
            "验证集与阈值校准：GMDCSA Subject 3 + Zenodo video split val。",
            "最终测试集：GMDCSA Subject 4，仅运行最终一次测试。",
            "阈值扫描只使用验证集，不使用 Subject4。",
            "Zenodo 数据按视频/subject 级别划分，不做窗口随机混合。",
        ],
    )

    doc.add_heading("3. 代码核心流程", level=1)
    add_numbered(
        doc,
        [
            "读取视频并切成固定窗口，窗口标签来自文件夹或视频路径中的类别信息。",
            "训练阶段先通过 CNN encoder 提取每帧空间特征，再由 projection 层映射到 hidden_dim。",
            "GRU 对窗口内时序特征进行建模，输出 mean / max / last 三种池化表示。",
            "SNN 分支从原始窗口帧中提取运动能量、人体重心、下落速度、形态变化等动态特征。",
            "SNN 分支通过 LIF-inspired 膜电位递推形成 spike 激活，再输出 mean / max / last 表示。",
            "CNN/GRU 表示与 SNN 表示拼接后进入融合 head，最终输出类别 logits。",
            "推理阶段输出逐帧 fall 概率，再用验证集选出的 fall-prob 阈值进行视频级判定。",
        ],
    )

    doc.add_heading("4. v4 的主要代码改动", level=1)
    add_bullets(
        doc,
        [
            "新增 --learnable-snn-dynamics：SNN 分支的 decay / threshold 改为每个通道可学习参数。",
            "增强 extract_temporal_event_features()：新增 center_drop_accel、relative_low_position、post_peak_hold 等跌倒瞬间特征。",
            "fall_event_like 更强调快速下落、姿态变宽、重心下降后的保持状态。",
            "run_lif_temporal_branch() 使用动态 decay / threshold 进行膜电位递推和 surrogate spike 激活。",
            "load_model_checkpoint() 已恢复 learnable_snn_dynamics 配置，保证训练和测试一致。",
        ],
    )

    doc.add_heading("5. SNN 分支技术逻辑", level=1)
    doc.add_paragraph("SNN 分支不是直接处理完整 RGB 图像，而是从视频窗口中提取动态序列特征。当前输入特征共 7 维：")
    add_bullets(
        doc,
        [
            "mass：归一化人体区域能量，近似反映人体区域占比。",
            "motion：相邻帧差分，反映运动强度。",
            "center_y：人体重心纵向位置。",
            "center_drop：重心向下移动量。",
            "soft_height：人体区域软高度。",
            "aspect：人体区域宽高比。",
            "fall_event_like：综合运动、下落、形态变化和跌倒后保持状态的事件分数。",
        ],
    )
    doc.add_paragraph("LIF-inspired 递推逻辑如下：")
    for line in [
        "current = snn_input(features)",
        "membrane = decay * membrane + current[:, t, :]",
        "spike = sigmoid((membrane - threshold) * surrogate_scale)",
        "membrane = membrane * (1.0 - spike.detach())",
    ]:
        p = doc.add_paragraph()
        run = p.add_run(line)
        set_font(run, "Consolas", 9.5)
    doc.add_paragraph(
        "其中 decay 表示膜电位保留比例，threshold 表示放电阈值。v4 中二者不再是全局固定值，"
        "而是通过 sigmoid 限制到安全范围后按通道学习：decay 位于 0.05-0.99，threshold 位于 0.05-1.25。"
    )

    doc.add_heading("6. 训练策略与关键参数", level=1)
    add_bullets(
        doc,
        [
            "初始化模型：来自 v3 checkpoint。",
            "前 2 个 epoch 冻结 CNN encoder / projection / GRU，只训练 SNN 分支和融合 head。",
            "第 3 个 epoch 起解除冻结，CNN、GRU、SNN 与融合 head 联合训练。",
            "teacher-student noise：开启，student noise mode 为 gaussian。",
            "Crossbar 器件动态：开启。",
            "crossbar_readout_noise_scale：0.25。",
            "hard_negative_normal_weight：1.10。",
            "hard_negative_min_percentile：0.88。",
            "hard_negative_score_gamma：2.00。",
            "normal_positive_penalty：0.00。",
            "external_labeled_stride：24。",
            "ADL context boost：Zenodo train/ADL，min_percentile=0.95，extra_copies=1。",
        ],
    )

    doc.add_heading("7. 训练过程摘要", level=1)
    add_bullets(
        doc,
        [
            "训练窗口：2376，其中 normal=1187，fall=1189。",
            "验证窗口：726，其中 normal=391，fall=335。",
            "ADL context boost：24 个视频，额外加入 44 个核心窗口。",
            "第 1-2 轮 CNN/GRU 冻结；第 3 轮开始联合训练。",
            "第 3 轮验证 balanced accuracy 最高，约 72.5%。",
            "第 8 轮触发 early stopping。",
        ],
    )

    doc.add_heading("8. 严格阈值扫描结果", level=1)
    doc.add_paragraph("阈值扫描只在 Subject3 + Zenodo val 上进行。推荐阈值为 0.55。")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells,
        ["threshold", "Acc", "Bal Acc", "ADL Spec", "Fall Recall", "Precision", "TP/TN/FP/FN"],
    ):
        cell.text = text
    for row in [
        ["0.55", "70.49%", "70.65%", "80.00%", "61.29%", "76.00%", "19/24/6/12"],
        ["0.60", "68.85%", "69.19%", "90.00%", "48.39%", "83.33%", "15/27/3/16"],
        ["0.50", "62.30%", "62.10%", "50.00%", "74.19%", "60.53%", "23/15/15/8"],
    ]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text

    doc.add_heading("9. Subject4 最终一次测试", level=1)
    doc.add_paragraph("Subject4 没有参与阈值或模型选择。本次仅使用验证集推荐阈值 0.55 进行最终测试。")
    add_metric_table(
        doc,
        [
            ("Accuracy", "64.86%"),
            ("Balanced Accuracy", "63.97%"),
            ("ADL Specificity", "75.00%"),
            ("Fall Recall", "52.94%"),
            ("Precision", "64.29%"),
            ("TP / TN / FP / FN", "9 / 15 / 5 / 8"),
            ("ADL 误报数", "5"),
            ("Fall 漏报数", "8"),
        ],
    )
    doc.add_paragraph("误报 ADL：05.mp4、06.mp4、11.mp4、12.mp4、13.mp4。")
    doc.add_paragraph("漏报 Fall：01.mp4、02.mp4、05.mp4、08.mp4、09.mp4、11.mp4、13.mp4、17.mp4。")

    doc.add_heading("10. 与历史最好结果的关系", level=1)
    doc.add_paragraph(
        "v4 虽然加入了更有论文意义的 SNN 动态分支，但当前 Subject4 结果仍未超过 7.4 无事件过滤版本。"
        "此前 7.4 无过滤器主结果约为 Accuracy 72.97%、ADL specificity 80.00%、Fall recall 64.71%、Precision 73.33%。"
        "这说明 SNN 分支的方向有意义，但当前参数和训练分配还没有充分解决 GMDCSA 跨 subject 的域差异。"
    )

    doc.add_heading("11. 创新点总结", level=1)
    add_bullets(
        doc,
        [
            "从纯 CNN/GRU 视频分类扩展为 CNN/GRU + LIF-inspired SNN temporal branch。",
            "SNN 输入使用运动能量、重心变化、姿态变化和 fall_event_like 动态序列，贴合跌倒瞬间快速状态变化。",
            "decay / threshold 设置为可学习通道参数，使动态响应不再完全依赖人工固定阈值。",
            "保留 Crossbar 器件动态和读出噪声，使模型结构与忆阻器工程背景保持关联。",
            "严格使用 Subject3 + 外部 val 做阈值校准，Subject4 独立测试，避免测试集污染。",
        ],
    )

    doc.add_heading("12. 当前问题与后续优化方向", level=1)
    add_bullets(
        doc,
        [
            "Fall recall 仍偏低，说明模型在独立 Subject4 上仍有偏保守倾向。",
            "误报集中在部分 ADL 视频，说明坐下、躺下、接近地面动作仍与跌倒混淆。",
            "SNN 分支已能表达动态变化，但还需要更合理的 Fall event positive mining 或损失权重设计。",
            "后续可尝试小范围扫描 SNN hidden_dim、surrogate_scale、freeze 轮数，以及 Fall event positive mining 权重。",
            "如果继续引入新增 ADL 数据，应控制权重，避免模型进一步倾向 ADL。",
        ],
    )

    doc.add_heading("13. 本次断点与漏点检查", level=1)
    add_bullets(
        doc,
        [
            "训练已完成，模型 checkpoint 和训练曲线均已生成。",
            "严格评估脚本已正常退出，无残留 my_yizu Python 进程。",
            "stderr 为空，未发现运行异常。",
            "阈值扫描文件、推荐阈值文件、Subject4 最终指标文件均已生成。",
            "v4 根目录原本缺少 README，本次已补充中文 README.md。",
            "本次流程未发现 Subject4 被用于调参或阈值选择。",
        ],
    )
    doc.add_paragraph("关键文件：")
    add_bullets(
        doc,
        [
            str(V4_DIR / "code" / "automatic_fall_detector.py"),
            str(V4_DIR / "models" / "automatic_fall_event_detector_cnn_snn_v4_learnable_snn.pth"),
            str(V4_DIR / "reports" / "v4_strict_validation_threshold_scan.csv"),
            str(V4_DIR / "reports" / "v4_subject4_final_metrics.json"),
        ],
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()


from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path(r'C:\Users\lin\Desktop\Ideal and non.docx')
d=Document(); s=d.sections[0]; s.top_margin=Inches(.65); s.bottom_margin=Inches(.65); s.left_margin=Inches(.75); s.right_margin=Inches(.75)
for n,z,c in [('Normal',9.5,None),('Title',19,'17365D'),('Heading 1',14,'1F4E79'),('Heading 2',11,'2F75B2')]:
 st=d.styles[n]; st.font.name='Microsoft YaHei'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); st.font.size=Pt(z)
 if c: st.font.color.rgb=RGBColor.from_string(c)
def shade(cell):
 x=OxmlElement('w:shd'); x.set(qn('w:fill'),'D9EAF7'); cell._tc.get_or_add_tcPr().append(x)
def table(head,rows):
 t=d.add_table(rows=1,cols=len(head)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,v in enumerate(head): t.rows[0].cells[i].text=str(v); shade(t.rows[0].cells[i])
 for row in rows:
  cs=t.add_row().cells
  for i,v in enumerate(row): cs[i].text=str(v)
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    for r in p.runs: r.font.name='Microsoft YaHei'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); r.font.size=Pt(8.5)
 d.add_paragraph()
def bullet(x): d.add_paragraph(x,style='List Bullet')
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Ideal and Non-Ideal Device Ablation Test Specification'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(23,54,93)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('忆阻器非理想特性与 Crossbar-CNN-GRU-SNN 消融测试文档').italic=True

d.add_heading('1. 测试目标',1); d.add_paragraph('本测试用于量化真实忆阻器器件特性对跌倒识别网络的影响，并区分性能变化来自电导状态、脉冲可塑性、噪声、循环波动还是长期漂移。实验结果必须同时报告 ADL 与 Fall，不能只报告 Accuracy。')
for x in ['第一层：比较 Ideal、Device-aware、Full non-ideal 三种总体条件。','第二层：在同一初始化和同一数据口径下，逐项加入一种非理想因素，寻找影响最大的变量。','第三层：对噪声强度、循环波动和漂移时间进行连续扫描，输出均值、标准差和性能曲线。']: bullet(x)

d.add_heading('2. 严格实验口径',1); table(['项目','固定设置'],[('工程基线',r'D:\ray\12.18-2\5.9new\crossbar+\CNN+SNN\v8，正式基线 v8 p006'),('模型',r'D:\ray\12.18-2\5.9new\crossbar+\CNN+SNN\v8\models\automatic_fall_event_detector_cnn_snn_v8_final_p006_penalty006.pth'),('数据','训练：GMDCSA Subject1-2 + Zenodo train；验证：GMDCSA Subject3 + Zenodo val'),('数据规模','train=3539（ADL 1758，Fall 1781）；val=1076（ADL 573，Fall 503）'),('输入','灰度图 64×64；时间窗口 16 帧；stride=16'),('评价','Accuracy、Balanced Accuracy、ADL specificity、Fall recall、Precision、混淆矩阵'),('重复','每个条件至少 5 次；记录 mean、std、min、max'),('Subject4','禁止用于参数选择；只在最终方案冻结后做一次独立测试')]); d.add_paragraph('每次运行日志必须确认：external_labeled_windows=3539、external_val_labeled_windows=1076、Loaded conductance data、Loaded device dynamics。任何一项不满足，该运行标记为 invalid。')

d.add_heading('3. 三组总体对比',1); table(['条件','必须包含的因素','目的'],[('Ideal','理想连续权重；无器件波动、噪声、漂移、LTP/LTD 非线性','获得软件模型上限性能'),('Device-aware','实测 HRS/LRS；实测 LTP/LTD；PPF/EPSC 时间响应；不加入额外随机噪声和长期漂移','评估器件本身特性影响'),('Full non-ideal','Device-aware + 读出噪声 + 循环波动 + 器件间差异 + 权重漂移','模拟完整硬件部署条件')]); d.add_paragraph('主图：分组柱状图。横轴为 Ideal、Device-aware、Full non-ideal；纵轴为 Accuracy、Balanced Accuracy、ADL specificity、Fall recall（单位 %）。每个条件使用 5 次运行均值，误差条为标准差。')

d.add_heading('4. 非理想因素逐项消融',1); table(['编号','新增因素','说明'],[('A0','Ideal','理想参考点'),('A1','+ HRS/LRS variation','高低阻状态分布和电导映射误差'),('A2','+ LTP/LTD nonlinearity','实测电导增强/减弱曲线及更新非线性'),('A3','+ PPF/EPSC dynamics','脉冲间隔和暂态电流响应；未进入网络时要明确说明'),('A4','+ readout noise','Crossbar 读出噪声，记录噪声类型和强度'),('A5','+ cycle-to-cycle variation','从实测循环分布采样 HRS/LRS、VSET/VRESET 或 conductance error'),('A6','+ device-to-device variation','不同阵元使用不同器件参数分布'),('A7','+ drift','按 retention/weight drift 更新权重'),('A8','Full non-ideal','A1-A7 全部开启')])

d.add_heading('5. 图表横纵坐标',1); table(['图','内容','横坐标','纵坐标'],[('Fig. A','三组总体指标','Ideal / Device-aware / Full non-ideal','四项指标（%），均值 ± 标准差'),('Fig. B','逐项非理想消融','A0-A8 或因素名称','Metric（%）或 ΔMetric（percentage points）'),('Fig. C','噪声鲁棒性','Noise level','Metric（%），均值 ± 标准差'),('Fig. D','漂移鲁棒性','Retention time / Drift step','Metric（%），均值 ± 标准差'),('Fig. E','循环波动','Cycle number','HRS、LRS、ON/OFF ratio 或 conductance'),('Fig. F','器件参数分布','VSET、VRESET、HRS、LRS 或 ON/OFF ratio','Probability density / Count'),('Fig. G','混淆矩阵','Predicted class：ADL / Fall','True class：ADL / Fall；颜色表示数量或比例')])

d.add_heading('6. 噪声与漂移扫描',1); d.add_paragraph('噪声和漂移必须分别比较：噪声是随机扰动，重点看均值和方差；漂移是时间相关变化，重点看性能随 retention time 的下降。'); table(['扫描','横坐标','纵坐标','重复'],[('Gaussian / readout noise','noise std 或 noise level','四项网络指标（%）','5 次'),('Salt-and-pepper / Poisson','corruption probability 或 intensity','四项网络指标（%）','5 次'),('Cycle variation','measured CV 或 variation scale','四项网络指标（%）','5 次'),('Weight drift','0、1、6、24、72 h 或 drift step','四项网络指标（%）','每点 5 次')]); d.add_paragraph('噪声曲线使用均值 ± 标准差；漂移曲线横轴为时间或漂移步数。只有仿真数据时，图注必须写明 model-based drift simulation。')

d.add_heading('7. 器件级测试与网络映射',1)
for x in ['I-V：横轴 Voltage，纵轴 Current，展示突变 SET/RESET 和多次循环曲线。','HRS/LRS：横轴 Cycle number 或 Device index，纵轴 Resistance，展示均值、分布和 ON/OFF ratio。','LTP/LTD：横轴 Pulse number，纵轴 Conductance，展示平均曲线和标准差阴影。','PPF：横轴 Pulse interval Δt，纵轴 PPF index=I2/I1。','EPSC：横轴 Time，纵轴 Current，记录 Ipeak、衰减时间常数和恢复过程。','网络映射：HRS/LRS 对应电导范围；LTP/LTD 对应权重更新；PPF/EPSC 对应时间动态；循环波动和漂移作为推理阶段扰动。']: bullet(x)

d.add_heading('8. CSV 输出要求',1); table(['文件名','必须字段'],[('overall_ideal_device_nonideal.csv','condition, seed, accuracy, balanced_accuracy, adl_specificity, fall_recall, precision'),('factor_ablation.csv','factor, enabled_factors, seed, accuracy, balanced_accuracy, adl_specificity, fall_recall'),('noise_robustness.csv','noise_type, noise_level, seed, accuracy, balanced_accuracy, adl_specificity, fall_recall'),('drift_robustness.csv','drift_model, drift_time, seed, accuracy, balanced_accuracy, adl_specificity, fall_recall'),('device_cycle_variation.csv','device_id, cycle, vset, vreset, hrs, lrs, on_off_ratio, success'),('device_plasticity.csv','device_id, sequence_id, pulse_index, conductance, mode, std_group')]); d.add_paragraph('所有 CSV 必须保留原始重复数据，不要只输出均值；均值和误差条由后处理脚本生成。')

d.add_heading('9. 最终验收标准',1)
for x in ['三组总体对比、逐项因素消融、噪声扫描和漂移扫描均有独立 CSV。','每个条件至少 5 次重复，报告 mean ± std。','所有实验使用相同 train/val 数据口径，且没有 Subject4 调参。','Ideal、Device-aware、Full non-ideal 的开关配置写入 JSON 或 README。','结果必须回答：真实器件造成多大性能代价、哪一种非理想因素影响最大、SNN 是否保留了对时间扰动的耐受性。','如果某因素尚未具备实测数据，必须标记为仿真因素，不能写成实测器件结果。']: bullet(x)

d.add_heading('10. 交接路径与约束',1)
for x in [r'项目基线：D:\ray\12.18-2\5.9new\crossbar+\CNN+SNN\v8',r'器件数据：D:\ray\12.18-2\data',r'Python：F:\ANACONDA\envs\my_yizu_3.10\python.exe',r'Skill：C:\Users\lin\.codex\skills\fall-crossbar-project\SKILL.md','不要修改 v8-final；新实验写入独立 models、reports、image 目录。','启动训练前先检查代码、数据口径和输出目录，不要自动重复启动已有进程。']: bullet(x)
d.add_paragraph('推荐论文逻辑：先用 Ideal 作为性能上限，再加入真实器件特性形成 Device-aware，最后加入噪声、循环波动和漂移形成 Full non-ideal；随后通过逐项消融和鲁棒性曲线解释性能变化来源。')
d.save(out); print(out)
